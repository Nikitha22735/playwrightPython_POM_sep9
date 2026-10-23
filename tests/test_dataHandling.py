import csv
import json
import os
import shutil
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
import pytest

DATA_DIR = Path(__file__).resolve().parents[1] / "testData"

# @pytest.mark.dh
def test_jsonhandling():
    with (DATA_DIR / "creds.json").open(encoding="utf-8") as data:
        creds = json.load(data)
        print(creds["NegitiveCreds"]["username"])


# @pytest.mark.dh
def test_csvHandling():
    with (DATA_DIR / "credentails.csv").open(
        newline="", encoding="utf-8"
    ) as data:
        creds = csv.DictReader(data)
        l1 = []
        for i in creds:
            l1.append(i)
        print(l1[2]["username"])


#pip install openpyxl

def test_handlingExcel():
    workbook = load_workbook(DATA_DIR / "sample_creds.xlsx")
    sheetData = workbook["Sheet2"]
    rows = [
        row
        for row in sheetData.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in row)
    ]

    print(rows[0][0])

# @pytest.mark.dh
def test_handlingExcelWrite(tmp_path):
    workbook_path = tmp_path / "sample_creds.xlsx"
    shutil.copyfile(DATA_DIR / "sample_creds.xlsx", workbook_path)
    workbook = load_workbook(workbook_path)
    sheetData = workbook["Sheet2"]
    # sheetData["A4"]="trpiura"
    # sheetData.delete_rows(2,sheetData.max_row)
    # sheetData.append(["tripra123","data123"])
    sheetData.insert_rows(2)
    sheetData["A2"]="helo"
    sheetData["B2"]="how"
    workbook.save(workbook_path)

# set sep9usName=testUser&&set sep9pw=testPassword&&pytest -m dh -s
# @pytest.mark.dh
def test_cli():
    username = os.getenv('sep9usName')
    password = os.getenv('sep9pw')
    print(username)
    print(password)

#pip install python-dotenv
# set file=.env.dev&&pytest -m dh -s
@pytest.mark.dh
def test_env():
    load_dotenv(os.getenv('file'))
    username = os.getenv('sep9usName_env1')
    password = os.getenv('sep9pw_env1')
    print(username)
    print(password)

